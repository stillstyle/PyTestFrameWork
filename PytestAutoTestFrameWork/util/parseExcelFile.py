from openpyxl import load_workbook
from config.conf import excelPath

class ParseExcel:
    def __init__(self):
        self.wk = load_workbook(excelPath)

    def getSheetByName(self, sheetName):
        return self.wk[sheetName]

    def getAllValuesOfSheet(self, sheet):
        maxRow = sheet.max_row
        maxCol = sheet.max_column
        allValues = []
        for row in range(2, maxRow + 1):
            rowValues = []
            for col in range(1, maxCol + 1):
                val = sheet.cell(row, col).value
                rowValues.append(val if val is not None else '')
            allValues.append(tuple(rowValues))
        return allValues